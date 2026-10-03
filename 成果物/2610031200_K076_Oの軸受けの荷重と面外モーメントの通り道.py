# K076 Oの継手で面外モーメントを受けるときの軸受けの荷重と、モーメントの通り道の選択肢（STEP4 4-3）— 直前の検討結果をスプレッドシートにまとめる
import subprocess
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
TS=subprocess.run(['date','+%y%m%d%H%M'],capture_output=True,text=True,env={'TZ':'Asia/Tokyo'}).stdout.strip()
FN='Yu Gothic'
f_in=Font(name=FN,color='0000FF'); f_n=Font(name=FN); f_b=Font(name=FN,bold=True); f_t=Font(name=FN,bold=True,size=13)
y=PatternFill('solid',start_color='FFFF00'); hd=PatternFill('solid',start_color='D9E1F2')
th=Side(style='thin',color='999999'); bd=Border(left=th,right=th,top=th,bottom=th); wr=Alignment(wrap_text=True,vertical='top')
wb=Workbook()
def widths(ws,ws_): 
    for i,w in enumerate(ws_): ws.column_dimensions[chr(65+i)].width=w
def hdr(ws,r,vals):
    for i,v in enumerate(vals):
        c=ws.cell(r,i+1,v); c.font=f_b; c.fill=hd; c.border=bd; c.alignment=wr
def cell(ws,r,c,v,font=f_n,fmt=None,fill=None):
    x=ws.cell(r,c,v); x.font=font; x.border=bd; x.alignment=wr
    if fmt: x.number_format=fmt
    if fill: x.fill=fill
    return x
# ---- 1 軸受けの荷重 ----
ws=wb.active; ws.title='軸受けの荷重'; widths(ws,[36,14,10,62])
ws['A1']='K076 Oで面外モーメントを受けるときの軸受け1個あたりの荷重'; ws['A1'].font=f_t
ws['A2']='青字＝入力値（変えると再計算）、黒字＝計算式。軸受け2個をx方向に間隔sで並べ、Mdを偶力で受ける。'; ws['A2'].font=f_n
hdr(ws,4,['入力','値','単位','根拠'])
cell(ws,5,1,'Oの面外モーメント Md（d軸まわり）'); cell(ws,5,2,82.3,f_in,'0.0'); cell(ws,5,3,'N·m'); cell(ws,5,4,'K066（2610021223版）使用位置・(0) Bは面外に固定・地震 KH=2.0・KV=1.0 の最大')
cell(ws,6,1,'Oの軸受けにかかる力（面内）'); cell(ws,6,2,1350,f_in,'#,##0'); cell(ws,6,3,'N'); cell(ws,6,4,'K067（2610021225版）リンク軸力の最大付近（収納・地震）。2個で等分と仮定')
cell(ws,7,1,'安全率'); cell(ws,7,2,4,f_in); cell(ws,7,3,''); cell(ws,7,4,'A04-007')
hdr(ws,9,['軸受けの間隔 s（x方向）','1個あたりの荷重','単位','注記'])
for i,s in enumerate((1.5,2.0,3.0,5.0,8.0)):
    r=10+i; cell(ws,r,1,s,f_in,'0.0" cm"')
    cell(ws,r,2,'=($B$5*$B$7/(A%d/100)+$B$6*$B$7/2)/1000'%r,fmt='0.0'); cell(ws,r,3,'kN')
    cell(ws,r,4,'梁面〜リンク面の間（左2.5・右2.0）に架台の板も入るため、今の配置で取れるのは1.5程度' if s<=2 else ('リンク面を内側へずらす必要あり（右はCSPの幅wRに響く）' if s>=5 else ''))
ws['A16']='注：Md・軸力はいずれも仮の前提（リンク断面・荷重係数など）に基づく計算値。軸受けの定格荷重との照合は未実施。'; ws['A16'].font=f_n; ws['A16'].alignment=wr; ws.merge_cells('A16:D16'); ws.row_dimensions[16].height=32
# ---- 2 申告 ----
w2=wb.create_sheet('申告'); widths(w2,[24,90])
w2['A1']='申告：K064〜K074の前提「Oは面外に剛」を部品で確かめていなかった'; w2['A1'].font=f_t
hdr(w2,3,['項目','内容'])
D=[('前提','K066以降の計算では、Oの継手が面外のモーメントMd（約80N·m）を受け止めるものとしていた。'),
   ('物理的な条件','モーメントを受けられるのは、軸受けを軸の方向（x）に間隔をあけて2個並べた場合だけ。'),
   ('今の配置','梁面からリンク面まで 左2.5・右2.0。架台の板も入るため、軸受けの間隔は1.5程度が限度 → 1個あたり20kN級となり、小さな軸受けでは持たない。'),
   ('影響範囲','依頼主案に限らず、今までの構成（K064〜K074の計算）すべてに関わる。'),
   ('原因','部品の当てはめを後回しにしていたため、気づくのが遅れた（Claude）。')]
for i,(a,b) in enumerate(D):
    cell(w2,4+i,1,a,f_b); cell(w2,4+i,2,b,fill=(y if a=='影響範囲' else None)); w2.row_dimensions[4+i].height=34
# ---- 3 選択肢 ----
w3=wb.create_sheet('モーメントの通り道'); widths(w3,[8,22,64,46,10])
w3['A1']='面外モーメントの通り道の選択肢'; w3['A1'].font=f_t
hdr(w3,3,['案','受ける場所','内容','気になる点','提案順'])
P=[('P1','O','軸受け2個をx方向に5〜8離して置き、リンク面を内側へずらす。','左は保持具の範囲に余裕あり。右はCSPの幅（wR）との兼ね合い。',2),
   ('P2','B','Oは軸受け1個のただのピンにし、面外のモーメントはBのキャリッジ（リニアガイドなどモーメントを受けられる直動部品）で受ける。レールは架台に多数のねじ（タップ止め、自前で可）で留め、モーメントをレールの長さに分散。','Bは使用時と収納時で位置が変わるため、梁のねじの配置もそれに合わせる必要。K065の(C)（Bを球面軸受）の逆の形。',1),
   ('P3','機構の外','P1・P2で足りない場合、使用位置でCSPを横から支えるなど、地震の横揺れを機構の外で受ける。','追加の部品と意匠への影響。','—')]
for i,row in enumerate(P):
    for j,v in enumerate(row): cell(w3,4+i,j+1,v)
    w3.row_dimensions[4+i].height=60
w3['A8']='Claudeの提案：P2を先に計算する（Oが小さなピンで済めば、依頼主案の「Oを下げてdを大きくする」配置や横柱のすき間が楽になるため）。'; w3['A8'].font=f_n; w3['A8'].alignment=wr; w3.merge_cells('A8:E8'); w3.row_dimensions[8].height=32
# ---- 4 依頼主の条件 ----
w4=wb.create_sheet('依頼主の条件'); widths(w4,[6,70,40])
w4['A1']='依頼主から示された条件（どの案でも使う）'; w4['A1'].font=f_t
hdr(w4,3,['No.','条件','状態'])
C=[('架台の板厚 8mm、梁のねじ 8本/片側','仮設定（依頼主）'),
   ('O・O′はhを下げ、dを大きくしてよい（現実的な軸受けを配置できるまで）','依頼主指示'),
   ('O′のまわりにも横柱を追加してよい','依頼主指示'),
   ('横柱 SF-20・20（20×20アルミ）：h=269・d=47、h=269・d=53、h=263・d=47','依頼主案（d=53の柱は位置を動かす必要あり：K075）'),
   ('全体のhが下がった分、収納時のCSPも下げる','依頼主指示'),
   ('Oの下の鉄板（T=10）はできるだけ小さく','依頼主指示'),
   ('右O′はd=68（左右の差18を保つ）','依頼主回答（「1」をQ1の答えとして解釈）')]
for i,(a,b) in enumerate(C):
    cell(w4,4+i,1,i+1); cell(w4,4+i,2,a); cell(w4,4+i,3,b); w4.row_dimensions[4+i].height=30
out='/mnt/user-data/outputs/%s_K076_Oの軸受けの荷重と面外モーメントの通り道.xlsx'%TS
wb.save(out); print(out)
