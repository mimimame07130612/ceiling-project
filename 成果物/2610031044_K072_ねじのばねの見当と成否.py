# K072 ねじのばね（ks）の見当と梁ねじの成否（STEP4 4-3）— 直前の検討結果をスプレッドシートにまとめる
import subprocess
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment
TS=subprocess.run(['date','+%y%m%d%H%M'],capture_output=True,text=True,env={'TZ':'Asia/Tokyo'}).stdout.strip()
FN='Yu Gothic'
f_in=Font(name=FN,color='0000FF'); f_n=Font(name=FN); f_b=Font(name=FN,bold=True); f_t=Font(name=FN,bold=True,size=13)
f_g=Font(name=FN,color='008000')
y=PatternFill('solid',start_color='FFFF00'); hd=PatternFill('solid',start_color='D9E1F2')
th=Side(style='thin',color='999999'); bd=Border(left=th,right=th,top=th,bottom=th)
wr=Alignment(wrap_text=True,vertical='top')
wb=Workbook()
def sheet(ws,widths):
    for i,w in enumerate(widths): ws.column_dimensions[chr(65+i)].width=w
def hdr(ws,r,vals):
    for i,v in enumerate(vals):
        c=ws.cell(r,i+1,v); c.font=f_b; c.fill=hd; c.border=bd; c.alignment=wr
def cell(ws,r,c,v,font=f_n,fmt=None,fill=None):
    x=ws.cell(r,c,v); x.font=font; x.border=bd; x.alignment=wr
    if fmt: x.number_format=fmt
    if fill: x.fill=fill
    return x

# ---- シート1 ks ----
ws=wb.active; ws.title='ksの見当'; sheet(ws,[34,14,12,60])
ws['A1']='K072 ねじの引抜き方向のばね ks の見当'; ws['A1'].font=f_t
ws['A2']='青字＝入力値（変えると再計算）、黒字＝計算式。黄色＝今後確かめたい前提。'; ws['A2'].font=f_n
hdr(ws,4,['入力','値','単位','根拠'])
rows=[('ねじ径 d',6,'mm','Xポイントビス よび径6.0（若井産業 木質構造用ねじカタログ 202504I改）'),
      ('ねじ部の有効長さ lef',30,'mm','DXP6100（全長100・ねじ長30）。梁の幅10.5（部屋.xlsx）を突き抜けない最長'),
      ('木材の密度 ρ（低）',330,'kg/m³','仮。梁の樹種・等級は未確認'),
      ('木材の密度 ρ（高）',420,'kg/m³','仮。梁の樹種・等級は未確認')]
for i,(a,b,c,d) in enumerate(rows):
    r=5+i; cell(ws,r,1,a); cell(ws,r,2,b,f_in,fill=(y if i>=2 else None)); cell(ws,r,3,c); cell(ws,r,4,d)
hdr(ws,10,['式','ks','単位','出典・注記'])
cell(ws,11,1,'Kser = 25·d·lef'); cell(ws,11,2,'=25*B5*B6',fmt='#,##0'); cell(ws,11,3,'N/mm'); cell(ws,11,4,'ETA-12/0062 などの形（KIT Karlsruher Berichte zum Ingenieurholzbau Bd.34 で紹介）')
cell(ws,12,1,'Kser = 780·d^0.2·lef^0.4'); cell(ws,12,2,'=780*B5^0.2*B6^0.4',fmt='#,##0'); cell(ws,12,3,'N/mm'); cell(ws,12,4,'ETA-11/0190 などの形（同上）')
cell(ws,13,1,'Kser = 234·ρ^0.2·d^0.2·lef^0.4（ρ低）'); cell(ws,13,2,'=234*B7^0.2*B5^0.2*B6^0.4',fmt='#,##0'); cell(ws,13,3,'N/mm'); cell(ws,13,4,'COST FP1402 STSM報告で比較されている式')
cell(ws,14,1,'Kser = 234·ρ^0.2·d^0.2·lef^0.4（ρ高）'); cell(ws,14,2,'=234*B8^0.2*B5^0.2*B6^0.4',fmt='#,##0'); cell(ws,14,3,'N/mm'); cell(ws,14,4,'同上')
cell(ws,16,1,'ks の見当（最小）',f_b); cell(ws,16,2,'=MIN(B11:B14)',f_b,'#,##0'); cell(ws,16,3,'N/mm')
cell(ws,17,1,'ks の見当（最大）',f_b); cell(ws,17,2,'=MAX(B11:B14)',f_b,'#,##0'); cell(ws,17,3,'N/mm')
ws['A19']='注：式はいずれも欧州の全ねじタイプの自己穿孔ねじ用。Xポイントビス（半ねじ）にそのまま当てはまるとは限らない。COST FP1402の報告では、モデル間の予測値の食い違いが大きいと指摘されている。'
ws['A19'].font=f_n; ws['A19'].alignment=wr; ws.merge_cells('A19:D19'); ws.row_dimensions[19].height=45

# ---- シート2 成否 ----
w2=wb.create_sheet('成否の見立て'); sheet(w2,[38,16,12,56])
w2['A1']='案(D3)の必要引抜き耐力とねじ長30の下限の比較（使用位置・地震 KH=2.0・右側）'; w2['A1'].font=f_t
hdr(w2,3,['ks（N/mm）','必要引抜き耐力（N）','','出典'])
for i,(k,v) in enumerate([(2000,1472),(5000,1791),(20000,2297)]):
    cell(w2,4+i,1,k,f_in,'#,##0'); cell(w2,4+i,2,v,f_in,'#,##0'); cell(w2,4+i,4,'K070（2610030957版）案(D3)、1本抜け×安全率4')
hdr(w2,8,['項目','値','単位','注記'])
cell(w2,9,1,'評価に使う ks'); x=cell(w2,9,2,"='ksの見当'!B11",f_g,'#,##0'); cell(w2,9,3,'N/mm'); cell(w2,9,4,'ksの見当シートの 25·d·lef（4式の最大）を参照。手入力に変えてもよい')
cell(w2,10,1,'必要引抜き耐力（ks=2,000〜5,000の直線補間）')
cell(w2,10,2,'=IF(B9<=A5,B4+(B5-B4)*(B9-A4)/(A5-A4),B5+(B6-B5)*(B9-A5)/(A6-A5))',fmt='#,##0'); cell(w2,10,3,'N'); cell(w2,10,4,'ksが5,000を超えるときは5,000〜20,000の間で補間')
cell(w2,11,1,'ねじ長30の引抜き 下限値'); cell(w2,11,2,1820,f_in,'#,##0'); cell(w2,11,3,'N'); cell(w2,11,4,'カタログ DXP6--0 ねじ長30：平均2.76kN・下限1.82kN（信頼水準75%の95%下側許容限界）。スギ機械等級E70以上、単調加力、6体。性能保証値ではない')
cell(w2,12,1,'余裕（下限値÷必要引抜き耐力）',f_b); cell(w2,12,2,'=B11/B10',f_b,'0.00"倍"'); cell(w2,12,3,'')
cell(w2,13,1,'判定',f_b); cell(w2,13,2,'=IF(B12>=1,"成立（余裕わずか）","不足")',f_b); cell(w2,13,4,'余裕1.0倍前後のため、前提の確認が必要')
cell(w2,15,1,'参考：ねじ長30の引抜き 平均値'); cell(w2,15,2,2760,f_in,'#,##0'); cell(w2,15,3,'N'); cell(w2,15,4,'同カタログ')
cell(w2,16,1,'参考：ねじ長45（DXP6130、全長130）'); cell(w2,16,2,'使用不可'); cell(w2,16,4,'全長130が梁の幅10.5を突き抜けるため。K063・K070での「ねじ長45なら余裕」は取り消し')

# ---- シート3 前提 ----
w3=wb.create_sheet('未確認の前提'); sheet(w3,[6,46,62])
w3['A1']='成否の見立てに重なっている、確かめていない前提'; w3['A1'].font=f_t
hdr(w3,3,['No.','前提','影響'])
P=[('ksの式は欧州の全ねじタイプ用','Xポイントビス（半ねじ）で ks が異なれば、必要引抜き耐力が変わる（ksが大きいほど厳しい）'),
   ('下限1.82kNはスギ機械等級E70以上での試験値','梁の樹種・等級が未確認。異なれば耐力が上下する'),
   ('押し込み側を「ねじのばね」で受けるモデル','実際は板が梁に当たって支える。回転の支点が板の縁に移り、引抜きが変わる可能性'),
   ('ねじの端あき・間隔の規定','木質構造用ねじの規定は未確認。丸鋼の例（縁距離4d・端距離7d）ならh=268（上面から2.0）は不可、h=257はぎりぎり'),
   ('左のO付近の上のねじと横材の重なり','K071で判明。ねじか横材の位置をずらす必要'),
   ('鋼板架台でのせん断値','カタログのせん断試験は木の側材。鋼板での値は未確認')]
for i,(a,b) in enumerate(P):
    cell(w3,4+i,1,i+1); cell(w3,4+i,2,a,fill=y); cell(w3,4+i,3,b); w3.row_dimensions[4+i].height=34

# ---- シート4 選択肢 ----
w4=wb.create_sheet('次の選択肢'); sheet(w4,[6,34,64,12])
w4['A1']='ここからの選択肢'; w4['A1'].font=f_t
hdr(w4,3,['No.','選択肢','内容','提案順'])
O=[('計算モデルを実際に近づける','押し込み側を「板が梁に当たる（押すだけ効くばね）」に直し、今の余裕の実態をつかむ',1),
   ('耐力側を上げる工夫を並べて比べる','O・Bのまわりのねじをさらに増やす、板を厚くする（梁面とリンク面のすき間 左2.5・右2.0の範囲で）など',2),
   ('荷重側の前提を監査する','KH=2.0（A04-006）・安全率4（A04-007）は前提。譲歩の余地の有無は依頼主の判断。Claudeから緩める提案はしない','—')]
for i,(a,b,c) in enumerate(O):
    cell(w4,4+i,1,i+1); cell(w4,4+i,2,a); cell(w4,4+i,3,b); cell(w4,4+i,4,c); w4.row_dimensions[4+i].height=34

# ---- シート5 出典 ----
w5=wb.create_sheet('出典'); sheet(w5,[40,90])
w5['A1']='出典'; w5['A1'].font=f_t
S=[('若井産業 木質構造用ねじカタログ 202504I改','https://www.wakaisangyo.co.jp/wp-content/themes/wakai_theme/images/top/2504ScrewForTimberStructureCatalog_web.pdf'),
   ('H.J. Blaß, Y. Steige: Steifigkeit axial beanspruchter Vollgewindeschrauben（KIT, Karlsruher Berichte Bd.34）','https://publikationen.bibliothek.kit.edu/1000085040/23845901'),
   ('COST Action FP1402 STSM報告（A. Ringhofer）Axially Loaded Self-Tapping Screws','https://webarchiv.typo3.tum.de/TUM/costfp1402/fileadmin/w00btl/www/All_Members/FP1402_STSM-report_Andreas-Ringhofer.pdf'),
   ('北海道立総合研究機構 林産試だより 2010年5月号（縁距離4d・端距離7d：丸鋼の例）','https://www.hro.or.jp/upload/49721/1100508.pdf'),
   ('K070 Oと Bのまわりにねじを集めたときの力','2610030957_K070_OとBのまわりにねじを集めたときの力.csv'),
   ('K071 架台とねじの配置','2610031002_K071_架台とねじの配置.svg'),
   ('部屋.xlsx（梁の幅10.5・高さ15.5）','KB')]
hdr(w5,3,['資料','URL・ファイル'])
for i,(a,b) in enumerate(S): cell(w5,4+i,1,a); cell(w5,4+i,2,b)
out='/mnt/user-data/outputs/%s_K072_ねじのばねの見当と成否.xlsx'%TS
wb.save(out); print(out)
